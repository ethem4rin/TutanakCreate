# -*- coding: utf-8 -*-
"""
Noksanlık/Aykırılık Bildirimi (.docx) üreteci.

Belge, örnek bildirimden türetilmiş ``sablon/noksanlik_bildirimi_sablon.docx``
dosyası üzerinden kurulur. Şablon; stilleri, numaralandırma tanımlarını,
üst/alt bilgileri, kenar boşluklarını ve sayfa düzenini olduğu gibi taşır.
Üretici yalnızca yer tutucuları doldurur ve ``{{NOKSANLIKLAR}}`` işaretçisinin
yerine noksanlık bloklarını yazar.

Her noksanlık bloğu üç paragraf kümesinden oluşur:

1. Noksanlık maddesi   — tutanaktan gelen metin, kalın, otomatik "1-, 2-, 3-"
                         numaralandırmalı liste paragrafı.
2. İlgili mevzuat      — seçilen her mevzuat maddesi için bir paragraf.
3. Müfettiş önerisi    — "Buna göre;" ile başlayan değerlendirme paragrafı.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from .metin import Parca, cozumle

KOK = Path(__file__).resolve().parent.parent
SABLON = KOK / "sablon" / "noksanlik_bildirimi_sablon.docx"

PUNTO = 11  # Örnek bildirimdeki gövde puntosu (w:sz = 22)

# Örnek bildirimden birebir alınan paragraf özellikleri
PPR_NOKSANLIK = (
    '<w:pStyle w:val="ListeParagraf"/>'
    '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="6"/></w:numPr>'
    '<w:tabs><w:tab w:val="left" w:pos="851"/><w:tab w:val="left" w:pos="1134"/></w:tabs>'
    '<w:spacing w:before="120" w:after="120" w:line="276" w:lineRule="auto"/>'
    '<w:ind w:left="0" w:firstLine="567"/>'
    '<w:contextualSpacing w:val="0"/>'
    '<w:jc w:val="both"/>'
    '<w:rPr><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr>'
)
PPR_DUZ = (
    '<w:tabs><w:tab w:val="left" w:pos="851"/><w:tab w:val="left" w:pos="1134"/></w:tabs>'
    '<w:spacing w:before="120" w:after="120" w:line="276" w:lineRule="auto"/>'
    '<w:jc w:val="both"/>'
    '<w:rPr><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr>'
)

YER_TUTUCU = re.compile(r"\{\{[A-Z0-9_]+\}\}")


@dataclass
class Noksanlik:
    """Bildirimde yer alacak tek bir noksanlık bloğu."""

    metin: str
    mevzuat: List[str] = field(default_factory=list)
    oneri: str = ""


@dataclass
class BildirimVerisi:
    sayi: str = ""
    tarih: str = ""
    unvan: str = ""
    adres: str = ""
    giris: str = ""
    bildirim_tarihi: str = ""
    bildirim_saati: str = ""
    sgk_sicil_no: str = ""
    eposta: str = ""
    mufettis_1: str = ""
    mufettis_2: str = ""
    teb_ad_soyad: str = ""
    teb_baba_ana: str = ""
    teb_dogum: str = ""
    teb_tckn: str = ""
    teb_gorev: str = ""
    teb_tarih: str = ""
    noksanliklar: List[Noksanlik] = field(default_factory=list)

    def yer_tutucular(self) -> Dict[str, str]:
        return {
            "{{SAYI}}": self.sayi,
            "{{TARIH}}": self.tarih,
            "{{UNVAN}}": self.unvan,
            "{{ADRES}}": self.adres,
            "{{GIRIS}}": self.giris,
            "{{BILDIRIM_TARIHI}}": self.bildirim_tarihi,
            "{{BILDIRIM_SAATI}}": self.bildirim_saati,
            "{{SGK_SICIL_NO}}": self.sgk_sicil_no,
            "{{EPOSTA}}": self.eposta,
            "{{MUFETTIS_1}}": self.mufettis_1,
            "{{MUFETTIS_2}}": self.mufettis_2,
            "{{TEB_AD_SOYAD}}": self.teb_ad_soyad,
            "{{TEB_BABA_ANA}}": self.teb_baba_ana,
            "{{TEB_DOGUM}}": self.teb_dogum,
            "{{TEB_TCKN}}": self.teb_tckn,
            "{{TEB_GOREV}}": self.teb_gorev,
            "{{TEB_TARIH}}": self.teb_tarih,
        }

    @classmethod
    def sozlukten(cls, veri: Dict[str, Any]) -> "BildirimVerisi":
        noksanliklar = [
            Noksanlik(
                metin=(n.get("metin") or "").strip(),
                mevzuat=[m for m in (n.get("mevzuat") or []) if (m or "").strip()],
                oneri=(n.get("oneri") or "").strip(),
            )
            for n in veri.get("noksanliklar", [])
            if (n.get("metin") or "").strip()
        ]
        alanlar = {k: (veri.get(k) or "") for k in cls.__dataclass_fields__ if k != "noksanliklar"}
        return cls(noksanliklar=noksanliklar, **alanlar)


# --------------------------------------------------------------------------- #
# Düşük seviyeli yardımcılar
# --------------------------------------------------------------------------- #

def _run_ekle(paragraf: Paragraph, parca: Parca) -> None:
    run = paragraf.add_run(parca.metin)
    run.font.size = Pt(PUNTO)
    if parca.kalin:
        run.font.bold = True
    if parca.italik:
        run.font.italic = True
    if parca.alti_cizili:
        run.font.underline = True
    # Örnek belgedeki gibi karmaşık yazı (szCs) puntosunu da sabitle.
    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn("w:szCs")) is None:
        rpr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="{PUNTO * 2}"/>'))


def _paragraf_kur(ppr_xml: str, parcalar: Iterable[Parca], parent) -> Paragraph:
    eleman = parse_xml(f'<w:p {nsdecls("w")}><w:pPr>{ppr_xml}</w:pPr></w:p>')
    paragraf = Paragraph(eleman, parent)
    for parca in parcalar:
        _run_ekle(paragraf, parca)
    return paragraf


def _paragraflari_gez(belge) -> Iterable[Paragraph]:
    """
    Gövde, tablo hücreleri ve üst/alt bilgilerdeki tüm paragrafları dolaşır.

    Üst/alt bilgiler ``section.header`` / ``section.footer`` yerine paket
    içindeki header/footer parçalarının XML'i üzerinden gezilir: bölüm
    nesnesi yalnızca o bölümün başvurduğu parçayı verir ve tablo içeriğini
    her zaman göstermez (alt bilgideki iletişim satırı bir tablo içindedir).
    """
    yield from belge.paragraphs
    for tablo in belge.tables:
        for satir in tablo.rows:
            for hucre in satir.cells:
                yield from hucre.paragraphs

    for parca in belge.part.package.iter_parts():
        ad = str(parca.partname)
        if "/header" not in ad and "/footer" not in ad:
            continue
        kok = getattr(parca, "element", None)
        if kok is None:
            continue
        for p in kok.iter(qn("w:p")):
            yield Paragraph(p, None)


def _yer_tutuculari_doldur(belge, degerler: Dict[str, str]) -> None:
    """
    Yer tutucuları run biçimlendirmesini bozmadan değiştirir.

    Şablonda her yer tutucu tek bir run içinde durduğu için run bazlı
    değiştirme yeterlidir; yine de paragraf metni üzerinden ikinci bir
    kontrol yapılır.
    """
    for paragraf in _paragraflari_gez(belge):
        if "{{" not in paragraf.text:
            continue
        for run in paragraf.runs:
            if "{{" not in run.text:
                continue
            yeni = run.text
            for anahtar, deger in degerler.items():
                if anahtar in yeni:
                    yeni = yeni.replace(anahtar, deger or "")
            # Doldurulmamış yer tutucu kalmasın.
            yeni = YER_TUTUCU.sub("", yeni)
            if yeni != run.text:
                run.text = yeni


def _isaretci_bul(belge, anahtar: str = "{{NOKSANLIKLAR}}") -> Optional[Paragraph]:
    for paragraf in belge.paragraphs:
        if anahtar in paragraf.text:
            return paragraf
    return None


def _bos_hucre(tablo, satir: int, sutun: int) -> None:
    hucre = tablo.rows[satir].cells[sutun]
    for paragraf in hucre.paragraphs:
        for run in list(paragraf.runs):
            run.text = ""


# --------------------------------------------------------------------------- #
# Ana üretim
# --------------------------------------------------------------------------- #

def noksanlik_bloklari(veri: BildirimVerisi, parent) -> List[Paragraph]:
    """Her noksanlık için madde + mevzuat + öneri paragraflarını üretir."""
    paragraflar: List[Paragraph] = []
    for noksanlik in veri.noksanliklar:
        paragraflar.append(
            _paragraf_kur(PPR_NOKSANLIK, [Parca(noksanlik.metin, kalin=True)], parent)
        )
        for madde in noksanlik.mevzuat:
            paragraflar.append(
                _paragraf_kur(PPR_DUZ, [Parca("\t")] + cozumle(madde.strip()), parent)
            )
        if noksanlik.oneri:
            paragraflar.append(
                _paragraf_kur(PPR_DUZ, [Parca("\t")] + cozumle(noksanlik.oneri.strip()), parent)
            )
    return paragraflar


def kur(veri: BildirimVerisi, sablon: str | Path = SABLON):
    """Bildirimi üretir ve ``docx.Document`` nesnesini döndürür."""
    sablon = Path(sablon)
    if not sablon.exists():
        raise FileNotFoundError(
            f"Şablon bulunamadı: {sablon}\n"
            "Önce `python tools/sablon_olustur.py` komutunu çalıştırın."
        )

    belge = docx.Document(str(sablon))

    isaretci = _isaretci_bul(belge)
    if isaretci is None:
        raise ValueError("Şablonda {{NOKSANLIKLAR}} işaretçisi bulunamadı.")

    for paragraf in noksanlik_bloklari(veri, isaretci._parent):
        isaretci._p.addprevious(paragraf._p)
    isaretci._p.getparent().remove(isaretci._p)

    # İmza tablosu: tek müfettiş varsa boş sütunun unvanı da silinir.
    if belge.tables:
        tablo = belge.tables[0]
        if not veri.mufettis_1.strip() and len(tablo.rows) >= 3:
            _bos_hucre(tablo, 2, 0)
        if not veri.mufettis_2.strip() and len(tablo.rows) >= 3:
            _bos_hucre(tablo, 2, 2)

    _yer_tutuculari_doldur(belge, veri.yer_tutucular())
    return belge


def bellege_uret(veri: BildirimVerisi, sablon: str | Path = SABLON) -> io.BytesIO:
    """Bildirimi diske hiç yazmadan bellekte üretir (sunucusuz ortam için)."""
    akis = io.BytesIO()
    kur(veri, sablon).save(akis)
    akis.seek(0)
    return akis


def olustur(veri: BildirimVerisi, cikti_yolu: str | Path, sablon: str | Path = SABLON) -> Path:
    """Bildirimi üretip ``cikti_yolu`` adresine kaydeder."""
    cikti_yolu = Path(cikti_yolu)
    cikti_yolu.parent.mkdir(parents=True, exist_ok=True)
    kur(veri, sablon).save(str(cikti_yolu))
    return cikti_yolu

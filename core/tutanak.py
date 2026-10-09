# -*- coding: utf-8 -*-
"""
Teftiş tutanağını (.docx) çözümler.

Tutanaktan iki şey çıkarılır:

1. **İşyeri ve teftiş künyesi** — noksanlık bildiriminin mavi (otomatik dolan)
   alanlarını besler: unvan, adres, görev yazısı bilgileri, teftiş tarihleri,
   müfettiş adları, SGK sicil numarası, işveren vekili.
2. **Noksanlık adayları** — tutanakta kalın (bold) punto ile yazılmış uzun
   paragraflar. Müfettiş arayüzde bunları onaylar ya da başka paragrafları
   elle işaretler.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional

import docx

# Kalınlık oranı bu eşiğin üzerindeyse paragraf noksanlık adayı sayılır.
KALINLIK_ESIGI = 0.60
# Bu uzunluğun altındaki kalın paragraflar başlıktır (örn. "V-) Çalışma Süreleri...").
ASGARI_NOKSANLIK_UZUNLUGU = 80
# Arayüzde listelenecek paragrafların asgari uzunluğu.
ASGARI_LISTELEME_UZUNLUGU = 40

# "I-", "VII-)", "IX-)" gibi roma rakamlı bölüm başlıkları
ROMA_BASLIK = re.compile(r"^\s*[IVXLCDM]+\s*[-)\.]")
# "5.1-)", "1.2-)", "3-" gibi madde numaraları
MADDE_NUMARASI = re.compile(r"^\s*\d+(?:\.\d+)*\s*[-–)\.]*\s*\)?\s*")


def tr_kucult(metin: str) -> str:
    """Türkçe'ye uygun küçük harfe çevirme."""
    return metin.replace("I", "ı").replace("İ", "i").lower()


@dataclass
class Paragraf:
    sira: int
    metin: str
    kalinlik_orani: float
    aday: bool
    baslik_mi: bool
    html: str


@dataclass
class Kunye:
    sicil_no: str = ""
    isyeri_unvani: str = ""
    isveren_vekili_ham: str = ""
    telefon: str = ""
    sgk_sicil_no: str = ""
    adres: str = ""
    tebligat_adresi: str = ""

    unvan: str = ""          # Bildirimde başlığa yazılacak tüzel unvan
    vekil_adi: str = ""
    vekil_gorevi: str = ""

    baskanlik: str = ""
    gorev_yazisi_tarihi: str = ""
    gorev_yazisi_sayisi: str = ""
    gorev_yazisi_eki: str = ""
    program_adi: str = ""
    teftis_tarihleri: str = ""
    tutanak_tarihi: str = ""

    mufettisler: List[str] = field(default_factory=list)
    giris_metni: str = ""


@dataclass
class TutanakSonucu:
    kunye: Kunye
    paragraflar: List[Paragraf]

    def sozluk(self) -> Dict[str, Any]:
        return {
            "kunye": asdict(self.kunye),
            "paragraflar": [asdict(p) for p in self.paragraflar],
        }


def _kalinlik_orani(paragraf) -> float:
    toplam = sum(len(r.text) for r in paragraf.runs)
    if not toplam:
        return 0.0
    kalin = sum(len(r.text) for r in paragraf.runs if r.font.bold)
    return kalin / toplam


def _html(paragraf) -> str:
    from html import escape

    parcalar = []
    for run in paragraf.runs:
        govde = escape(run.text)
        if run.font.bold:
            govde = f"<strong>{govde}</strong>"
        if run.font.italic:
            govde = f"<em>{govde}</em>"
        parcalar.append(govde)
    return "".join(parcalar)


def madde_numarasini_ayikla(metin: str) -> str:
    """Noksanlık metninin başındaki "5.1-)" gibi numarayı siler."""
    return MADDE_NUMARASI.sub("", metin, count=1).strip()


def _bilgi_tablosu(belge) -> Dict[str, str]:
    """Tutanağın başındaki künye tablosunu etiket -> değer sözlüğüne çevirir."""
    bilgiler: Dict[str, str] = {}
    for tablo in belge.tables:
        if len(tablo.columns) < 2:
            continue
        for satir in tablo.rows:
            hucreler = [h.text.strip() for h in satir.cells]
            if len(hucreler) >= 3 and hucreler[1] in (":", "："):
                etiket, deger = hucreler[0], hucreler[2]
            elif len(hucreler) == 2:
                etiket, deger = hucreler[0], hucreler[1]
            else:
                continue
            etiket = tr_kucult(etiket.replace("\t", " ")).strip(" :")
            if etiket and deger and deger != "-":
                bilgiler.setdefault(etiket, deger)
    return bilgiler


def _bul(bilgiler: Dict[str, str], *anahtarlar: str) -> str:
    for anahtar in anahtarlar:
        aranan = tr_kucult(anahtar)
        for etiket, deger in bilgiler.items():
            if aranan in etiket:
                return deger
    return ""


def _mufettisler(belge) -> List[str]:
    """İmza tablosundan "İş Müfettişi" unvanlı kişilerin adlarını toplar."""
    adlar: List[str] = []
    for tablo in belge.tables:
        if len(tablo.rows) < 2:
            continue
        satirlar = [[h.text.strip() for h in satir.cells] for satir in tablo.rows]
        unvan_satiri = None
        for sira, satir in enumerate(satirlar[1:], start=1):
            if any("müfettiş" in tr_kucult(h) for h in satir):
                unvan_satiri = sira
                break
        if unvan_satiri is None:
            continue
        for sutun, unvan in enumerate(satirlar[unvan_satiri]):
            if "müfettiş" in tr_kucult(unvan) and sutun < len(satirlar[0]):
                ad = satirlar[0][sutun].strip()
                if ad and ad not in adlar:
                    adlar.append(ad)
    return adlar


def _vekil_ayikla(ham: str) -> Dict[str, str]:
    """
    "Örnek Sanayi ve Tic. A.Ş./ İnsan Kaynakları Müdür Yrd. – Ad Soyad"
    biçimindeki metni unvan / görev / ad olarak ayırır.
    """
    sonuc = {"unvan": "", "gorev": "", "ad": ""}
    ham = (ham or "").strip()
    if not ham:
        return sonuc

    if "/" in ham:
        unvan, kalan = ham.split("/", 1)
    else:
        unvan, kalan = ham, ""
    sonuc["unvan"] = unvan.strip(" ./")

    kalan = kalan.strip()
    for ayrac in ("–", "—", "-"):
        if ayrac in kalan:
            gorev, ad = kalan.split(ayrac, 1)
            sonuc["gorev"] = gorev.strip()
            sonuc["ad"] = ad.strip()
            break
    else:
        sonuc["gorev"] = kalan
    return sonuc


def _giris_bilgileri(belge, kunye: Kunye) -> None:
    """Tutanağın giriş paragrafından görev yazısı / program / tarih bilgilerini çıkarır."""
    giris = ""
    for paragraf in belge.paragraphs:
        if "görev yazısı" in tr_kucult(paragraf.text):
            giris = paragraf.text
            break
    if not giris:
        return

    eslesme = re.search(
        r"(?P<baskanlik>Çalışma ve Sosyal Güvenlik Bakanlığı\s+.+?)\s+Başkanlığı",
        giris,
    )
    if eslesme:
        kunye.baskanlik = re.sub(r"\s+", " ", eslesme.group("baskanlik")).strip()

    eslesme = re.search(
        r"(?P<tarih>\d{1,2}[./]\d{1,2}[./]\d{4})\s*tarihli\s*ve\s*"
        r"(?P<sayi>[^\s(]+)\s*(?:\(\s*(?P<ek>[^)]*?)\s*\))?",
        giris,
    )
    if eslesme:
        kunye.gorev_yazisi_tarihi = eslesme.group("tarih").replace(".", "/")
        kunye.gorev_yazisi_sayisi = eslesme.group("sayi")
        kunye.gorev_yazisi_eki = (eslesme.group("ek") or "").strip()

    eslesme = re.search(r"[“\"]([^”\"]{10,})[”\"]", giris)
    if eslesme:
        kunye.program_adi = eslesme.group(1).strip()

    eslesme = re.search(r"(?P<tarihler>\d[\d.,/\s]*\d)\s*tarih(?:ler)?inde", giris)
    if eslesme:
        kunye.teftis_tarihleri = re.sub(r"\s+", "", eslesme.group("tarihler"))


def _tutanak_tarihi(belge) -> str:
    for paragraf in reversed(belge.paragraphs):
        if "nüsha olarak düzenlendi" in tr_kucult(paragraf.text) or "düzenlendi" in tr_kucult(
            paragraf.text
        ):
            eslesme = re.search(r"(\d{1,2}[./]\d{1,2}[./]\d{4})", paragraf.text)
            if eslesme:
                return eslesme.group(1).replace("/", ".")
    return ""


def giris_metni_uret(kunye: Kunye) -> str:
    """Noksanlık bildiriminin giriş paragrafını künyeden kurar."""
    baskanlik = kunye.baskanlik or "Çalışma ve Sosyal Güvenlik Bakanlığı"
    sayi = kunye.gorev_yazisi_sayisi
    if kunye.gorev_yazisi_eki:
        sayi = f"{sayi} ({kunye.gorev_yazisi_eki})"
    return (
        f"{baskanlik} Başkanlığının {kunye.gorev_yazisi_tarihi} tarihli ve {sayi} sayılı "
        f"görev yazısı gereği yukarıda unvanı ve adresi belirtilen işyerinde "
        f"“{kunye.program_adi}” kapsamında çalışma mevzuatı yönünden denetim yapmak amacıyla "
        f"{kunye.teftis_tarihleri} tarihlerinde yapılan programlı teftişte tespit edilen "
        f"noksanlık ve aykırılıklar aşağıda belirtilmiştir."
    )


def coz(kaynak) -> TutanakSonucu:
    """
    Tutanağı çözümleyip künye ve paragraf listesini döndürür.

    ``kaynak`` bir dosya yolu ya da dosya benzeri nesne (ör. yüklenen dosyanın
    akışı) olabilir; sunucusuz ortamda tutanak diske yazılmadan okunur.
    """
    belge = docx.Document(str(kaynak) if isinstance(kaynak, (str, Path)) else kaynak)

    bilgiler = _bilgi_tablosu(belge)
    kunye = Kunye(
        sicil_no=_bul(bilgiler, "sicil no.", "sicil no"),
        isyeri_unvani=_bul(bilgiler, "işyerinin unvanı", "unvan"),
        isveren_vekili_ham=_bul(bilgiler, "işveren/işveren vekili", "işveren vekili", "işveren"),
        telefon=_bul(bilgiler, "telefon"),
        sgk_sicil_no=_bul(bilgiler, "sgk sicil"),
        adres=_bul(bilgiler, "işyerinin adresi", "adresi"),
        tebligat_adresi=_bul(bilgiler, "tebligat adresi"),
    )
    kunye.tebligat_adresi = kunye.tebligat_adresi or kunye.adres
    kunye.adres = kunye.adres or kunye.tebligat_adresi

    vekil = _vekil_ayikla(kunye.isveren_vekili_ham)
    kunye.unvan = vekil["unvan"] or kunye.isyeri_unvani
    kunye.vekil_adi = vekil["ad"]
    kunye.vekil_gorevi = vekil["gorev"]

    kunye.mufettisler = _mufettisler(belge)
    kunye.tutanak_tarihi = _tutanak_tarihi(belge)
    _giris_bilgileri(belge, kunye)
    kunye.giris_metni = giris_metni_uret(kunye)

    paragraflar: List[Paragraf] = []
    for sira, paragraf in enumerate(belge.paragraphs):
        metin = paragraf.text.strip()
        if len(metin) < ASGARI_LISTELEME_UZUNLUGU:
            continue
        oran = _kalinlik_orani(paragraf)
        baslik_mi = bool(ROMA_BASLIK.match(metin))
        aday = (
            oran >= KALINLIK_ESIGI
            and len(metin) >= ASGARI_NOKSANLIK_UZUNLUGU
            and not baslik_mi
        )
        paragraflar.append(
            Paragraf(
                sira=sira,
                metin=metin,
                kalinlik_orani=round(oran, 3),
                aday=aday,
                baslik_mi=baslik_mi,
                html=_html(paragraf),
            )
        )

    return TutanakSonucu(kunye=kunye, paragraflar=paragraflar)

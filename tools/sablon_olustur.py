# -*- coding: utf-8 -*-
"""
Örnek "Noksanlık Bildirimi" Word belgesinden, program tarafından kullanılacak
şablonu (sablon/noksanlik_bildirimi_sablon.docx) üretir.

Örnek belgedeki tüm biçimlendirme (yazı tipi, punto, satır aralığı, girintiler,
üst bilgi/alt bilgi, kenar boşlukları, stiller, numaralandırma tanımları)
olduğu gibi korunur; yalnızca değişken metinler yer tutucularla değiştirilir
ve noksanlık blokları tek bir işaretçi paragrafa indirgenir.

Kullanım:
    python tools/sablon_olustur.py
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn

KOK = Path(__file__).resolve().parent.parent
KAYNAK = KOK / "ornek" / "Vatan Noksanlik Bildirimi.docx"
HEDEF = KOK / "sablon" / "noksanlik_bildirimi_sablon.docx"

EPOSTA = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]{2,}")

# Örnek belgedeki paragraf indeksleri (çözümleme sonucu sabittir)
IDX_SAYI = 1
IDX_UNVAN = 5
IDX_ADRES = 6
IDX_GIRIS = 8
IDX_NOKSANLIK_BAS = 9
IDX_NOKSANLIK_SON = 19  # dahil
IDX_KAPANIS2 = 22
IDX_TEB_BAS = 30
IDX_TEB_SON = 35


def rpr_kopyala(kaynak_run, hedef_run) -> None:
    """Bir run'ın karakter biçimlendirmesini (rPr) aynen kopyalar."""
    kaynak_rpr = kaynak_run._r.find(qn("w:rPr"))
    if kaynak_rpr is None:
        return
    mevcut = hedef_run._r.find(qn("w:rPr"))
    if mevcut is not None:
        hedef_run._r.remove(mevcut)
    hedef_run._r.insert(0, copy.deepcopy(kaynak_rpr))


def paragrafi_yeniden_yaz(paragraf, parcalar) -> None:
    """
    Paragrafın pPr'ını koruyarak run'larını yeniden kurar.

    parcalar: [(metin, kalin_mi_ya_da_None), ...]
        kalin_mi None ise, paragrafın ilk run'ının biçimi aynen kullanılır.
    """
    ornek_run = paragraf.runs[0] if paragraf.runs else None
    ornek_r = copy.deepcopy(ornek_run._r) if ornek_run is not None else None

    for run in list(paragraf.runs):
        run._r.getparent().remove(run._r)

    for metin, kalin in parcalar:
        yeni = paragraf.add_run(metin)
        if ornek_r is not None:
            rpr = ornek_r.find(qn("w:rPr"))
            if rpr is not None:
                mevcut = yeni._r.find(qn("w:rPr"))
                if mevcut is not None:
                    yeni._r.remove(mevcut)
                yeni._r.insert(0, copy.deepcopy(rpr))
        if kalin is not None:
            yeni.font.bold = kalin


def paragrafi_sil(paragraf) -> None:
    paragraf._p.getparent().remove(paragraf._p)


def hucre_metnini_degistir(hucre, yeni_metin: str) -> None:
    """Hücredeki ilk paragrafın biçimini koruyarak metni değiştirir."""
    paragraflar = hucre.paragraphs
    if not paragraflar:
        return
    paragrafi_yeniden_yaz(paragraflar[0], [(yeni_metin, None)])
    for fazla in paragraflar[1:]:
        if fazla.text.strip():
            paragrafi_yeniden_yaz(fazla, [("", None)])


def _ust_alt_paragraflari(belge):
    """
    Üst/alt bilgilerdeki tüm ``w:p`` düğümlerini döndürür.

    ``section.header`` / ``section.footer`` yalnızca o bölümün başvurduğu
    parçayı verir ve tablo içeriğini her zaman göstermez; bu yüzden paket
    içindeki header/footer parçalarının XML'i doğrudan dolaşılır.
    """
    for parca in belge.part.package.iter_parts():
        ad = str(parca.partname)
        if "/header" not in ad and "/footer" not in ad:
            continue
        kok = getattr(parca, "element", None)
        if kok is None:
            continue
        yield from kok.iter(qn("w:p"))


def kisisel_verileri_temizle(belge) -> None:
    """
    Örnek belgeden gelen kişisel verileri şablondan siler.

    * Alt bilgideki müfettiş e-posta adresi ``{{EPOSTA}}`` yer tutucusuna çevrilir
      (adres, telefon ve faks Grup Başkanlığına ait kurumsal bilgilerdir, kalır).
    * Belgeyi oluşturan/son kaydeden kişi gibi meta veriler temizlenir.
    """
    for paragraf in _ust_alt_paragraflari(belge):
        metin_dugumleri = list(paragraf.iter(qn("w:t")))
        metin = "".join(d.text or "" for d in metin_dugumleri)
        if not EPOSTA.search(metin):
            continue
        metin_dugumleri[0].text = EPOSTA.sub("{{EPOSTA}}", metin)
        metin_dugumleri[0].set(qn("xml:space"), "preserve")
        for fazla in metin_dugumleri[1:]:
            fazla.text = ""

    ozellikler = belge.core_properties
    ozellikler.author = "Müfettiş Programı"
    ozellikler.last_modified_by = ""
    ozellikler.comments = ""
    ozellikler.category = ""
    ozellikler.keywords = ""
    ozellikler.subject = ""
    ozellikler.title = ""
    ozellikler.revision = 1


def main() -> int:
    if not KAYNAK.exists():
        print(f"HATA: Örnek belge bulunamadı: {KAYNAK}", file=sys.stderr)
        return 1

    belge = docx.Document(str(KAYNAK))
    p = belge.paragraphs

    # 1) Sayı / Tarih satırı
    paragrafi_yeniden_yaz(
        p[IDX_SAYI],
        [("Sayı  :", True), (" {{SAYI}}", False), ("\t" * 8, False), ("{{TARIH}}", False)],
    )

    # 2) İşyeri unvanı ve adresi (ortalanmış, kalın)
    paragrafi_yeniden_yaz(p[IDX_UNVAN], [("{{UNVAN}}", None)])
    paragrafi_yeniden_yaz(p[IDX_ADRES], [("{{ADRES}}", None)])

    # 3) Giriş paragrafı
    paragrafi_yeniden_yaz(p[IDX_GIRIS], [("{{GIRIS}}", None)])

    # 4) Noksanlık blokları -> tek işaretçi paragraf
    #    İlk paragrafın biçimi bize gerekmiyor; düz bir işaretçi bırakıyoruz.
    isaretci = p[IDX_NOKSANLIK_BAS]
    for i in range(IDX_NOKSANLIK_BAS + 1, IDX_NOKSANLIK_SON + 1):
        paragrafi_sil(p[i])
    ppr = isaretci._p.find(qn("w:pPr"))
    if ppr is not None:
        isaretci._p.remove(ppr)
    paragrafi_yeniden_yaz(isaretci, [("{{NOKSANLIKLAR}}", None)])

    # 5) Kapanış paragrafındaki bildirim tarih/saati
    # `p` silme işleminden önce alındığı için indeksler kaymaz.
    kapanis = p[IDX_KAPANIS2]
    sabit_metin = (
        "Tarafınıza tebliğ edilen iş bu bildirim, işyerinizde tespit edilen noksanlık ve "
        "aykırılıkların tarafınızdan giderileceğine ilişkin beyanınız üzerine düzenlenmiş olup "
        "işyerinizde tespit edilen noksanlık ve aykırılıkların verilen süre sonunda "
        "giderilebilmesi için müfettişliğimizce teftişe ara verilmiştir. Teftiş, verilen süre "
        "sonunda yapılacak denetim sonucunda düzenlenecek tutanakla tamamlanacaktır. Teftiş "
        "Raporu, müteakip denetim sonucunda düzenlenecek tutanağın ardından yazılacaktır. "
        "Buna göre; işverenliğinizin talebi ve müfettişliğimizce bu talebin uygun görülmesi "
        "üzerine, bu noksanlık ve aykırılıkların giderilmesini, belirtilen süre sonunda teftiş "
        "tamamlanacağından "
    )
    kalan_metin = (
        " işyerinde noksanlık ve aykırılıkların giderildiğine ilişkin defter, kayıt ve "
        "belgelerin hazır bulundurulmasını, aksi halde ayrıca hakkınızda 4857 sayılı İş "
        "Kanunu’nun 92/II ve 107/a maddeleri uyarınca cezai işlem yapılacağının bilinmesini "
        "tebliğen rica ederiz."
    )
    paragrafi_yeniden_yaz(
        kapanis,
        [
            (sabit_metin, False),
            ("{{BILDIRIM_TARIHI}} tarihinde saat {{BILDIRIM_SAATI}}’da", True),
            (kalan_metin, False),
        ],
    )

    # 6) Tebellüğ Eden alanları
    teb_alanlari = [
        ("Adı Soyadı", "\t\t: ", "{{TEB_AD_SOYAD}}"),
        ("Baba/Ana Adı", "\t\t: ", "{{TEB_BABA_ANA}}"),
        ("Doğum Yeri/Tarihi", "\t: ", "{{TEB_DOGUM}}"),
        ("T.C. Kimlik Numarası", ": ", "{{TEB_TCKN}}"),
        ("Görevi", "\t\t\t: ", "{{TEB_GOREV}}"),
        ("Tarih Ve İmza/Kaşe", "\t: ", "{{TEB_TARIH}}"),
    ]
    for sira, (etiket, ayrac, yer_tutucu) in enumerate(teb_alanlari):
        hedef = p[IDX_TEB_BAS + sira]
        paragrafi_yeniden_yaz(hedef, [(etiket + ayrac, None), (yer_tutucu, None)])

    # 7) Müfettiş imza tablosu
    tablo = belge.tables[0]
    hucre_metnini_degistir(tablo.rows[0].cells[0], "{{MUFETTIS_1}}")
    hucre_metnini_degistir(tablo.rows[0].cells[2], "{{MUFETTIS_2}}")

    # 8) Üst bilgideki işyeri sicil numarası
    for bolum in belge.sections:
        for ust in (bolum.header, bolum.first_page_header, bolum.even_page_header):
            for paragraf in ust.paragraphs:
                if "Sicil No" in paragraf.text:
                    paragrafi_yeniden_yaz(
                        paragraf, [("İşyeri Sicil No.\t: ", None), ("{{SGK_SICIL_NO}}", None)]
                    )

    # 9) Kişisel verileri temizle
    kisisel_verileri_temizle(belge)

    HEDEF.parent.mkdir(parents=True, exist_ok=True)
    belge.save(str(HEDEF))

    print(f"Şablon oluşturuldu: {HEDEF}")
    print("--- Şablon paragrafları ---")
    for i, paragraf in enumerate(docx.Document(str(HEDEF)).paragraphs):
        print(f"[{i}] {paragraf.style.name}: {paragraf.text[:110]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

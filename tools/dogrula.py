# -*- coding: utf-8 -*-
"""
Üretilen bildirimi örnek bildirimle karşılaştırır.

Örnek tutanaktan künye çıkarılır, örnek bildirimdeki üç noksanlık bloğu
kütüphaneden beslenir ve üretilen belge; paragraf metni, biçimlendirme
(kalın/italik/altı çizili/punto), paragraf stili, hizalama, girinti ve
aralık değerleri bakımından örnekle karşılaştırılır.

Kullanım:
    python tools/dogrula.py
"""
from __future__ import annotations

import sys
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK))

import docx  # noqa: E402

from core import belge as belge_modulu  # noqa: E402
from core import kutuphane, mevzuat, tutanak  # noqa: E402

ORNEK_TUTANAK = KOK / "ornek" / "VATAN Tutanak.docx"
ORNEK_BILDIRIM = KOK / "ornek" / "Vatan Noksanlik Bildirimi.docx"
CIKTI = KOK / "calisma" / "cikti" / "_dogrulama.docx"


def madde(kimlik: str) -> str:
    kayit = mevzuat.madde_getir(kimlik)
    if kayit is None or not kayit["paragraf"]:
        raise SystemExit(f"Mevzuat maddesi yok ya da metni boş: {kimlik}")
    return kayit["paragraf"]


def oneri(kimlik: str) -> str:
    for kayit in kutuphane.yukle("oneri"):
        if kayit["id"] == kimlik:
            return kayit["metin"]
    raise SystemExit(f"Öneri kaydı yok: {kimlik}")


def paragraf_ozeti(belge):
    satirlar = []
    for paragraf in belge.paragraphs:
        bicim = paragraf.paragraph_format
        satirlar.append(
            {
                "stil": paragraf.style.name,
                "hiza": str(paragraf.alignment),
                "girinti": (
                    str(bicim.left_indent),
                    str(bicim.right_indent),
                    str(bicim.first_line_indent),
                ),
                "aralik": (
                    str(bicim.space_before),
                    str(bicim.space_after),
                    str(bicim.line_spacing),
                ),
                "metin": paragraf.text,
                "runlar": [
                    (r.text, bool(r.font.bold), bool(r.font.italic), bool(r.font.underline),
                     str(r.font.size))
                    for r in paragraf.runs
                    if r.text
                ],
            }
        )
    return satirlar


def main() -> int:
    sonuc = tutanak.coz(ORNEK_TUTANAK)
    kunye = sonuc.kunye

    adaylar = [p for p in sonuc.paragraflar if p.aday]
    # Örnek bildirimdeki 3. madde (6.1) tutanakta kalın yazılmadığı için
    # otomatik aday listesinde yoktur; müfettişin elle eklediği durumu taklit
    # etmek üzere burada doğrudan sıra numarasıyla alınıyor.
    ucuncu = next(p for p in sonuc.paragraflar if p.metin.startswith("6.1-)"))

    metinler = [tutanak.madde_numarasini_ayikla(p.metin) for p in adaylar]
    metinler.append(tutanak.madde_numarasini_ayikla(ucuncu.metin))

    veri = belge_modulu.BildirimVerisi(
        sayi="34283/05-34282/05",
        tarih="07.08.2026",
        unvan=kunye.unvan,
        adres=kunye.adres,
        giris=kunye.giris_metni,
        bildirim_tarihi="05/10/2026",
        bildirim_saati="10.00",
        sgk_sicil_no=kunye.sgk_sicil_no,
        mufettis_1="Örnek MÜFETTİŞ",
        mufettis_2="Örnek MÜFETTİŞ",
        teb_ad_soyad=kunye.vekil_adi,
        teb_baba_ana="Örnek / Örnek",
        teb_dogum="Örnek İlçe / 01.01.1980",
        teb_tckn="00000000000",
        teb_gorev=kunye.vekil_gorevi,
        teb_tarih=kunye.tutanak_tarihi,
        noksanliklar=[
            belge_modulu.Noksanlik(
                metin=metinler[0],
                mevzuat=[madde("kanun|4857|63"), madde("yonetmelik|fazla-calisma|5")],
                oneri=oneri("oneri-calisma-suresi"),
            ),
            belge_modulu.Noksanlik(
                metin=metinler[1],
                mevzuat=[madde("kanun|4857|46")],
                oneri=oneri("oneri-hafta-tatili"),
            ),
            belge_modulu.Noksanlik(
                metin=metinler[2],
                mevzuat=[madde("kanun|4857|47")],
                oneri=oneri("oneri-ubgt"),
            ),
        ],
    )

    belge_modulu.olustur(veri, CIKTI)

    uretilen = paragraf_ozeti(docx.Document(str(CIKTI)))
    ornek = paragraf_ozeti(docx.Document(str(ORNEK_BILDIRIM)))

    rapor = KOK / "calisma" / "cikti" / "_dogrulama_raporu.txt"
    satirlar = [f"Üretilen: {len(uretilen)} paragraf | Örnek: {len(ornek)} paragraf", ""]
    for sira in range(max(len(uretilen), len(ornek))):
        u = uretilen[sira] if sira < len(uretilen) else None
        o = ornek[sira] if sira < len(ornek) else None
        ayni = u == o
        metin_ayni = (u or {}).get("metin") == (o or {}).get("metin")
        isaret = "AYNI " if ayni else ("METIN=" if metin_ayni else "FARK ")
        satirlar.append(f"--- [{sira}] {isaret}")
        if not ayni:
            satirlar.append(f"  ÜRETİLEN: {u}")
            satirlar.append(f"  ÖRNEK   : {o}")
        else:
            satirlar.append(f"  {(u or {}).get('metin','')[:120]}")
    rapor.write_text("\n".join(satirlar), encoding="utf-8")

    farkli = sum(
        1
        for sira in range(max(len(uretilen), len(ornek)))
        if (uretilen[sira] if sira < len(uretilen) else None)
        != (ornek[sira] if sira < len(ornek) else None)
    )
    print(f"Üretilen belge : {CIKTI}")
    print(f"Karşılaştırma  : {rapor}")
    print(f"Paragraf sayısı: üretilen={len(uretilen)} örnek={len(ornek)}")
    print(f"Farklı paragraf: {farkli}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
